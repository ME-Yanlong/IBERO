"""演讲稿绘图组件；命令行默认只增补 V4，不改写 V3 原页。"""

from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
# 文稿工具隔离在产物目录，不升级仿真环境中的 Pillow / 其它依赖。
sys.path.insert(0, str(ROOT / "artifacts/pptx_tools"))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SOURCE = ROOT / "IBERO_工业级双臂具身智能操作基准平台演讲PPT_v3.pptx"
OUTPUT = ROOT / "IBERO_工业级双臂具身智能操作基准平台演讲PPT_v4.pptx"
VIDEO = ROOT / "artifacts/demo_videos_20260928"
OUT = ROOT / "artifacts/presentation_v4"
BG, CARD, BORDER = "071D2B", "10364C", "2B5C70"
WHITE, MUTED, TEAL, BLUE, GOLD = "F0F6FC", "A8C3D0", "59D7CE", "43A6FF", "FFB64D"
FONT = "Microsoft YaHei"
prs = Presentation(SOURCE)
original = list(prs.slides)
order = []
entries = []


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=CARD, stroke=None, rounded=False):
    shape = slide.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE if rounded else SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.13
    return shape


def text(slide, value, x, y, w, h, size=17, fill=WHITE, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "V4: " + value[:40].replace("\n", " / ")
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fill)
        p.space_after = Pt(3)
        p.line_spacing = 1.13
        if align is not None:
            p.alignment = align
        for run in p.runs:
            ea = OxmlElement("a:ea")
            ea.set("typeface", FONT)
            run._r.get_or_add_rPr().append(ea)
    return shape


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def badge(slide, value, fill=TEAL):
    text(slide, value, 9.4, 0.30, 2.85, 0.24, 10, fill, True, PP_ALIGN.RIGHT)


def new_slide(title, section, status="Core-0.1 · 已实现原型", subtitle="", source=""):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    text(slide, section, 0.66, 0.31, 8.4, 0.24, 10, TEAL, True)
    badge(slide, status, GOLD if "规划" in status or "目标" in status else TEAL)
    text(slide, title, 0.66, 0.67, 12.0, 0.59, 26, WHITE, True)
    box(slide, 0.66, 1.37, 12.0, 0.008, BORDER)
    if subtitle:
        text(slide, subtitle, 0.66, 1.60, 12.0, 0.7, 17, WHITE)
    text(slide, source or "IBERO · Core-0.1 / 从工业约束到可验证的仿真机制", 0.66, 7.13, 11.9, 0.22, 8.3, MUTED)
    order.append(slide)
    entries.append({"slide": len(order), "title": title, "kind": "new", "source": source})
    return slide


def replace(slide, old, new):
    """保留原 V3 文本框位置及首个 run 的样式，不改变原文件。"""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text == old:
            p = shape.text_frame.paragraphs[0]
            style = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            r = p.add_run()
            r.text = new
            if style is not None:
                r._r.insert(0, style)
            for extra in list(shape.text_frame._txBody)[2:]:
                if extra.tag.endswith("}p") and extra is not p._p:
                    shape.text_frame._txBody.remove(extra)
            return True
    raise ValueError(f"Cannot find original text: {old}")


def preserve(index, status=None):
    slide = original[index - 1]
    order.append(slide)
    entries.append({"slide": len(order), "title": f"V3 第 {index} 页（保留并校订）", "kind": "adapted", "original_slide": index})
    if status:
        badge(slide, status, GOLD)
    return slide


def card(slide, title, body, x, y, w, h, accent=TEAL, body_size=17):
    box(slide, x, y, w, h, CARD, BORDER, True)
    box(slide, x + 0.18, y + 0.24, 0.045, 0.29, accent)
    text(slide, title, x + 0.35, y + 0.21, w - 0.55, 0.42, 20, accent, True)
    text(slide, body, x + 0.25, y + 0.79, w - 0.50, h - 0.88, body_size)


def ribbon(slide, value, y=6.55, fill=TEAL):
    box(slide, 0.66, y, 12.0, 0.39, "12394A", None, True)
    text(slide, value, 0.83, y + 0.065, 11.66, 0.27, 12, fill, True)


def flow(slide, labels, y, w=2.55, gap=0.47, x=0.66, h=1.25):
    for i, (title, detail) in enumerate(labels):
        left = x + i * (w + gap)
        box(slide, left, y, w, h, CARD, BORDER, True)
        text(slide, title, left + 0.13, y + 0.15, w - 0.26, 0.38, 18, TEAL, True, PP_ALIGN.CENTER)
        text(slide, detail, left + 0.13, y + 0.64, w - 0.26, 0.50, 12.5, MUTED, False, PP_ALIGN.CENTER)
        if i + 1 < len(labels):
            text(slide, "→", left + w + 0.10, y + 0.41, gap - 0.08, 0.38, 22, GOLD, True)


def video_slide(stem, title, key, takeaways, caution, mechanism, seek):
    video = VIDEO / f"{stem}.mp4"
    meta = json.loads(video.with_suffix(".json").read_text(encoding="utf-8"))
    if not meta["full_decode_verified"] or not meta["source_case_passed"]:
        raise ValueError(f"Unverified video: {video}")
    if hashlib.sha256(video.read_bytes()).hexdigest() != meta["video_sha256"]:
        raise ValueError(f"Video digest mismatch: {video}")
    slide = new_slide(title, f"DEMO / {key}", subtitle="", source=f"真实轨迹离线回放 · 仿真时间 1× · 源码 {meta['source_commit'][:7]} / seed {meta['seed']} · 非实时性能证明")
    movie = slide.shapes.add_movie(str(video), Inches(0.66), Inches(1.51), Inches(7.06), Inches(5.5303),
                                   poster_frame_image=str(video.with_suffix(".poster.png")), mime_type="video/mp4")
    movie.name = f"IBERO embedded video / {stem}"
    box(slide, 8.02, 1.51, 4.64, 5.53, CARD, BORDER, True)
    text(slide, f"{meta['video_duration_s']:.0f} 秒  /  单击画面播放", 8.26, 1.78, 4.1, 0.37, 19, TEAL, True)
    text(slide, "观察重点", 8.26, 2.39, 4.1, 0.34, 17, BLUE, True)
    text(slide, "\n".join(f"{i+1}. {v}" for i, v in enumerate(takeaways)), 8.26, 2.91, 4.05, 1.89, 16)
    text(slide, "实现依据", 8.26, 4.89, 4.1, 0.3, 16, TEAL, True)
    text(slide, mechanism, 8.26, 5.30, 4.08, 0.65, 14.5)
    text(slide, caution, 8.26, 6.24, 4.04, 0.59, 12.2, GOLD)
    notes(slide, f"演讲提示：{seek}\n视频已内嵌，请在 PowerPoint 放映模式点击画面播放；无音轨。末帧为封面，并非尚未播放时的实时状态。\n来源：{video.name}；原始回合 seed={meta['seed']}；源码={meta['source_commit']}；仿真区间={meta['simulation_start_s']:.3f}–{meta['simulation_end_s']:.3f}s；视频={meta['video_duration_s']:.2f}s（含片头片尾静帧）。\n来源哈希={meta['source_hash']}；视频 SHA256={meta['video_sha256']}。\n这是版本对应的成功单回合，不替代最终 S9 同源完整回归，不宣称材料已实测标定。{caution}")
    entries[-1].update(kind="video", video=video.name, sha256=meta["video_sha256"], duration_s=meta["video_duration_s"], source_commit=meta["source_commit"])
    return slide


def _legacy_expanded_main():
    OUT.mkdir(parents=True, exist_ok=True)
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    s = preserve(1)
    replace(s, "面向高校教授与工业界研发负责人", "Core-0.1 实现与七段仿真演示 · V4")
    replace(s, "IBERO · 2026", "IBERO · V4")
    notes(s, "V4 在原建设愿景之上补充 Core-0.1 实现、机制逻辑和七段内嵌视频。统一使用 IBERO 名称。所有工业级、高保真、开放生态表述为建设目标，不表示已完成工业认证、实机部署或最终全量验收。")
    s = preserve(2)
    replace(s, "现有机器人基准更多回答“能不能抓起来”，工业现场需要回答“能否在约束下长期可靠地完成工艺”。", "面向工业工艺，基准需要进一步覆盖双臂协同、复杂接触，以及可审计的质量与安全约束。")
    replace(s, "桌面单臂、低接触复杂度", "通用操作与任务泛化")
    replace(s, "难以覆盖双臂协同与高精度接触", "仍需更系统的工业接触评测")
    replace(s, "高混合、小批量、旧件非标", "小批量、多品种、旧件非标")
    # 原模板这行紧贴框宽；适度缩小字号，避免中英文标点挤出卡片。
    for shape in s.shapes:
        if shape.has_text_frame and shape.text == "小批量、多品种、旧件非标":
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(15)
    s = preserve(3, "平台使命 / 建设目标")
    replace(s, "平台方负责", "平台方建设目标")
    s = preserve(4, "总体蓝图 / 分阶段交付")
    replace(s, "IBERO的价值不在于单个模型，而在于把“研究问题”组织成可复用的工程闭环。", "先完成 Core 中可验证的接触与材料机制，再逐步接入数据、训练、社区工具和榜单。")
    replace(s, "双臂同步遥操作 · LeRobot / HDF5 / Zarr 数据 · 场景校验 · CI/CD 接入", "规划：同步遥操作 · LeRobot / HDF5 / Zarr · 社区脚手架与 CI 接入")
    replace(s, "双臂物理动力学 · 力位混合控制 · 公共工作空间 · ROS 2 / Gymnasium 接口", "当前：MuJoCo + 接触/材料/加工原型；ROS 2 与统一数据接口后续建设")
    s = preserve(5, "参与价值 / 预期收益")
    replace(s, "在数字孪生中评估换线、拆装和再制造可行性；沉淀企业 CAD 与工艺库", "先标定关键材料与接触，再逐步评估工艺；沉淀脱敏 CAD 与工艺库")
    replace(s, "用 ibero-cli 接入场景，参与榜单与工业双臂挑战赛", "当前贡献场景与回归证据；CLI、榜单和赛事按路线图建设")
    s = preserve(6, "分层设计 / 上层逐步接入")
    replace(s, "物理仿真 · 双臂动力学 · 力位控制 · 统一接口", "MuJoCo · 机器人/材料/机构/加工 · 控制/任务/审计")
    notes(s, "Core 是当前实施重点，不能把设计中的神经符号中间件、形式化安全、排行榜或 ROS 2 / LeRobot 接口说成已经交付。当前具体运行类按场景工作，不提前构造通用多引擎框架。")

    s = new_slide("Core-0.1：先把工业交互的物理闭环做实", "CORE / CURRENT IMPLEMENTATION", subtitle="当前成果：5 类任务环境、7 段演示；核心差异在接触、材料响应和可审计的状态变化。")
    card(s, "物理底座", "MuJoCo 3.8.1\nG1 上半身 + 场景指定工具\n真实接触、关节动力学与积分", 0.66, 2.40, 5.80, 1.86, BLUE, 16)
    card(s, "三类对象 / 机制", "一维 Flex 线束\n被动弹性舌片 + 卡扣接触\n体素工件 + 过程级切削载荷", 6.86, 2.40, 5.80, 1.86, TEAL, 16)
    card(s, "策略与验证", "脚本基线驱动仿真机器人\n任务终止 / 越界审计 / 证据报告\n轨迹身份与材料事件可回放", 0.66, 4.43, 5.80, 1.86, TEAL, 16)
    card(s, "明确不混淆", "不是宇树实机 SDK 下发\n尚非工业材料实测数字孪生\n未交付学习策略榜单或形式化安全", 6.86, 4.43, 5.80, 1.86, GOLD, 16)
    ribbon(s, "先保证“为什么能完成、为什么会失败”可解释，再扩展训练接口与更大 benchmark。")
    notes(s, "当前演示使用仓库内 G1 / Robotiq 模型与 MuJoCo，不连接实体机器人或厂商驱动。五类环境是 cable_handover、cable_stretch、latch_release、harness_unplug、plate_milling；后者提供三种目标形状。七段视频不等于七个已完成规模化基准套件。")

    s = new_slide("一次仿真如何组成：配置 → 编译 → 闭环 → 证据", "CORE / RUNTIME ARCHITECTURE", subtitle="场景描述对象与目标；Core 负责把它们变成可重复执行、可检查的物理实验。", source="代码索引：core/scene_loader.py · scene_compiler.py · milling_builder.py · milling_runtime.py")
    flow(s, [("场景菜谱", "YAML + task_spec.py"), ("校验与编译", "机器人 / 材料 / 工装"), ("MuJoCo 运行", "控制、载荷、接触、积分"), ("任务与审计", "成功 / 失败 / 证据")], 2.53, w=2.62, gap=0.50)
    card(s, "每个控制周期", "读取当前状态 → 脚本 / 控制器产生命令\n物理子步：材料力 / 加工力 → 动力学积分\n读取真实姿态、接触、传感器与材料状态", 0.66, 4.02, 7.25, 2.06, BLUE, 17)
    card(s, "记录与显示分离", "Trace / StockTrace 存状态与事件\n审阅器恢复轨迹，不续跑策略\nMP4 按仿真时钟离线渲染", 8.21, 4.02, 4.45, 2.06, TEAL, 15)
    ribbon(s, "加工 builder 与 runtime 已拆分；当前按真实对象复用组件，不提前铺开通用多引擎框架。")
    notes(s, "示意图表达组件关系，不意味着所有环境已统一 Gymnasium API。场景编译创建模型和稳定句柄；运行层处理时钟、输入、载荷合成、积分与审计。加工分离构建与运行；线束/卡扣仍保留各自具体环境。")

    s = new_slide("哪些由物理求解，哪些由模型和任务规则定义？", "CORE / PHYSICS VERSUS RULES", subtitle="采用“刚体接触求解 + 显式材料 / 加工模型 + 只读任务判定”的组合，而不是成功动画。")
    card(s, "引擎求解", "关节和刚体动力学\n接触 / 摩擦 / 约束反力\n机器人受力与真实运动", 0.66, 2.55, 3.80, 2.66, BLUE, 18)
    card(s, "可标定的近似模型", "Flex 轴向 + 中心线弯曲\n舌片离散弹性梁\n体素去除 + 平均切削载荷", 4.76, 2.55, 3.80, 2.66, TEAL, 18)
    card(s, "任务与安全规则", "基于状态判断解锁 / 放置\n检测力、位姿、侵入等越界\n记录失败，不伪造完成", 8.86, 2.55, 3.80, 2.66, GOLD, 18)
    text(s, "接触使物体运动；材料模型决定其响应；task_spec.py 判断结果。", 0.80, 5.62, 11.8, 0.45, 22, WHITE, True)
    ribbon(s, "“未触发阈值”不等于“真实材料绝对无损”；“检测到碰撞”不等于形式化防碰撞屏障。")
    notes(s, "锁止状态机只观察卡扣的几何/接触结果，不通过布尔开关释放焊接约束。切削采用工程过程模型与几何去除，不宣称第一性原理的金属塑性/断裂求解。材料参数目前仍 provisional。")

    s = new_slide("P2 / P3 / P4：谁负责什么，参数放在哪里？", "CORE / CONTRIBUTOR BOUNDARIES", subtitle="配置归场景，物理实现归组件；任务可以组合对象，但不应绕开物理直接写成功。")
    card(s, "P2 机械师", "robots/\nG1 裁剪、夹爪 / 按压工具 / 主轴\n关节、安装位姿、TCP 与传感器\n工具替换移除重叠的原手部外观", 0.66, 2.48, 3.80, 3.16, BLUE, 16)
    card(s, "P3 材料专家", "materials/ · mechanisms/ · processes/\n线束、弹性梁、卡扣、工件\n物性 → 力学响应 / 材料事件\n不拥有某个任务的成功逻辑", 4.76, 2.48, 3.80, 3.16, TEAL, 15)
    card(s, "P4 菜谱作者", "scenes/<任务>/\nscene_config.yaml：对象与工况\nconstraints.yaml：约束与阈值\ntask_spec.py：阶段及终止条件", 8.86, 2.48, 3.80, 3.16, GOLD, 15)
    ribbon(s, "长度 / 外径 / 线密度 / 弯曲半径 / 端头 / 刚度 / 摩擦等来自场景配置；来源与 seed 一起记录。")
    notes(s, "约束阈值与材料物性区分：例如最小弯曲半径可以是报警指标，不表示自动模拟真实破坏。G1 原生手选项仅保留当前资产的静态几何，不等于已经实现可控制的灵巧手。ROS 2、LeRobot 和遥操作接口仍按后续阶段建设。")

    s = new_slide("线束：一维 Flex + 轴向弹簧阻尼 + 弯曲能", "MATERIAL / CABLE", subtitle="线束不是一根跟着末端移动的装饰线：它拥有自由节点、质量、接触和可计算的张力。", source="源码：materials/cable.py · mechanics.py · harness.py；配置：scenes/cable_handover 与 cable_stretch")
    flow(s, [("场景物性", "长度 / 线密度 / 刚度"), ("自由节点与边", "MuJoCo 1D Flex"), ("拉伸与弯曲", "轴向力 + 能量负梯度"), ("状态观测", "张力 / 曲率 / 能量")], 2.50, w=2.62, gap=0.50)
    card(s, "端头连接 ≠ 抓取捷径", "材料端点与端头永久连接，表示真实装配关系。\n夹爪与端头仍依赖接触和摩擦；\n交接时不通过临时 weld 把物体粘到手上。", 0.66, 4.04, 7.14, 2.12, BLUE, 17)
    card(s, "当前建模边界", "无扭转的离散中心线模型\n不是多层护套 / 导体实体应力\n不是自动断裂或材料寿命预测", 8.10, 4.04, 4.56, 2.12, GOLD, 16)
    ribbon(s, "当前线束默认已是 Flex；早期 capsule-chain 路线不应再写成现行默认。")
    notes(s, "与先前阶段契约相比，当前 scene_config.yaml 明确 representation: flex；旧 capsule_chain 别名已禁用。轴向刚度/阻尼由原生 Flex 边求解；弯曲力来自离散能量 E=EI/(2h)Σ|t[i+1]-t[i]|² 的负梯度。此处的一维 Flex 与卡扣三维实体 Flex 候选是不同对象/模型，不能混为一谈。")
    video_slide("01_cable_handover", "Demo 01｜双臂线束端头交接与放置", "CABLE HANDOVER", ["接收夹爪建立接触夹持", "原夹爪释放，支撑关系转移", "末端送到接料区并放开"], "刚性端头交接，不宣称裸线束抓取。", "Flex 线束 + 夹爪真实接触；任务读取支撑与目标距离。", "完整观看约 10 秒，重点看两手释放先后与最终工件支撑。")
    video_slide("02_cable_stretch", "Demo 02｜双臂拉伸与张力反馈", "CABLE TENSION", ["两端夹持使线束由松弛变紧", "控制器读取实际材料张力", "在目标范围内保持并接受审计"], "目标阈值是任务约束，不是实测破坏载荷。", "轴向边弹簧/阻尼 + 弯曲力；两端控制命令改变实际运动。", "约 13 秒。看线束从弯曲变直、张力读数变化与保持阶段，不把目标张力当外加恒力。")

    s = new_slide("卡扣：必须先让扣齿产生真实几何净空", "MECHANISM / ELASTIC SNAP LATCH", subtitle="插头能否拔出由扣齿与肩部的接触决定；按压使被动舌片变形，状态标签只是结果。", source="源码：materials/elastic_beam.py · mechanisms/snap_latch.py · control/latch.py")
    flow(s, [("未按压", "扣齿—肩部接触阻挡"), ("工具按压", "接触力弯曲舌片"), ("形成净空", "实际几何越过肩部"), ("退出与回弹", "拉出 / 松开 / 弹性恢复")], 2.47, w=2.62, gap=0.50)
    card(s, "舌片如何变形", "分段刚体 + 被动转角弹簧 / 阻尼\n刚度来自材料 E、截面 I 与段长 h\n普通段 kθ = EI/h；根部采用半单元修正", 0.66, 4.04, 7.18, 2.09, TEAL, 17)
    card(s, "如何证明不是开关", "观察净空、退出距离与接触力\n加入不按压硬拔等反事实测试\n不关锁止接触，不运行中删约束", 8.14, 4.04, 4.52, 2.09, BLUE, 16)
    ribbon(s, "当前为平面弹性梁近似；三维实体 Flex 候选不是默认，塑性损伤 / 疲劳 / 断裂未完成。")
    notes(s, "elastic_beam 使用 8 段（场景可配）被动 hinge，k=EI/h，根部2EI/h，阻尼k*tau。snap_latch 的扣齿和肩部是实际碰撞几何，动态插头有自由关节。状态机不执行解锁。台架反事实测试与机器人任务不能互相替代；本演讲录像为机器人成功单例。")
    video_slide("03_latch_release", "Demo 03｜按压弹性卡扣并拔出", "LATCH RELEASE", ["按压工具接触并压弯舌片", "净空由负变正，再开始退出", "撤离后舌片回弹，插头已分离"], "示例材料参数，未标定具体商品连接器。", "弹性梁 + 扣齿/肩部碰撞；夹持依赖真实接触。", "约 14 秒。结合右下角净空和接触力曲线讲清按压与解锁的先后因果。")
    video_slide("04_harness_unplug", "Demo 04｜带线束的解锁、拔出与放置", "HARNESS UNPLUG", ["复用同一个卡扣与按压机制", "线束随插头运动并产生张力", "放到接料盘后释放夹爪"], "接料盘投放，不宣称无落差精密装配。", "卡扣 + Flex 线束 + 任务链；端头质量只计一次。", "约 18 秒。展示可组合的工业对象：卡扣机制不变，新增线束动力学与投放阶段。")

    s = new_slide("铣削：载荷反馈与材料去除是同一个闭环", "PROCESS / MATERIAL REMOVAL", subtitle="采用固定体素工件 + 三轴过程级平均载荷；孔槽既改变可见形状，也改变实际碰撞占据。", source="源码：materials/stock*.py · processes/milling*.py · core/loads.py · core/milling_runtime.py")
    flow(s, [("读取真实运动", "刀具位姿 / 进给 / 转速"), ("计算切削载荷", "啮合 + 过程系数"), ("提交材料事件", "占据 / 碰撞 / 版本"), ("积分与反馈", "力 / 扭矩回到机器人")], 2.50, w=2.62, gap=0.50)
    card(s, "孔是怎样形成的", "刀具扫掠体找出待去除体素；合规才提交。\n同一事件更新材料账本、碰撞与显示状态。\n按目标检查孔槽，探针验证去除区域可通过。", 0.66, 4.02, 7.20, 2.14, TEAL, 17)
    card(s, "力不是旁路曲线", "力 / 扭矩进入统一载荷汇总\n机器人与主轴受载荷影响\n不以截断载荷伪造安全成功", 8.16, 4.02, 4.50, 2.14, BLUE, 16)
    ribbon(s, "不是金属塑性断裂 FEM；飞屑仅可视化；当前左臂加工、右臂停放，不是双臂协同铣削。")
    notes(s, "MillingProcess 读取实际运动，计算啮合和平均切削载荷；预测扫掠几何并做有效性检查，事务提交去除事件，载荷写入工具及固定夹具反作用账目，再由 runtime mj_step。当前模型有隐式载荷响应处理，不是根据目标路径预先挖洞。力超限/无有效啮合等工况可判失效。工件固定，不支持自由余料、五轴、多物理场切屑形成。")
    video_slide("05_milling_through_hole", "Demo 05｜铣削通孔：从实体到贯通", "THROUGH HOLE", ["主轴启动，工具向实体进给", "观察材料剖面逐层去除", "贯通后保留孔形与载荷记录"], "飞屑仅可视化；右臂安全停放。", "真实主轴/机器人运动 + 体素碰撞去除 + 过程载荷。", "约 46 秒；可在演讲中快进到切入阶段。俯视为便于观察仅隐藏工具显示，不改变工具物理。")
    video_slide("06_milling_slot", "Demo 06｜铣削直槽：路径改变加工形状", "SLOT MILLING", ["先下刀，再沿槽方向进给", "扫掠去除形成连续槽形", "观察转速、载荷与材料体积"], "当前展示浅槽，不把它标成贯通孔。", "同一材料/加工内核，场景目标与控制路径不同。", "约 53 秒。强调不是换一张工件贴图：体素事件与碰撞形状同步，最后有实际槽底。")
    video_slide("07_milling_pocket", "Demo 07｜铣削浅腔：二维区域逐步去除", "POCKET MILLING", ["连续轨迹逐步扩大去除区域", "浅腔底部仍保留实体材料", "停止主轴后核对最终腔形"], "示例过程系数，未宣称真实钢材切削精度。", "可重复的材料去除事件 + 形状指标 + 物理状态轨迹。", "约 86 秒。对照左下俯视与右下截面，区分浅腔与贯穿洞；完整视频可会后观看。")

    s = new_slide("演示成功之后，还需要什么证据？", "VALIDATION / EVIDENCE AND LIMITS", subtitle="单回合视频回答“能否发生”；工程回归与实测标定回答“是否可信、稳定、可迁移”。")
    card(s, "已经具备的检查层", "组件级：梁响应 / 材料账本 / 接触台架\n任务级：正确动作与错误动作、变体 / 种子\n回放级：源码 / 场景 / 资产身份与材料哈希", 0.66, 2.52, 7.30, 2.10, BLUE, 17)
    card(s, "本次视频能证明的范围", "七条已完成的成功轨迹\n对应版本、seed 和原始报告可追溯\n不代替最终 S9 同源全量验收", 8.26, 2.52, 4.40, 2.10, TEAL, 15.5)
    card(s, "仍需补齐", "真实材料曲线与参数标定 → 留出工况验证\n性能优化与完整回归 → 后续实机 / ROS 2 / LeRobot / 学习策略接入", 0.66, 4.78, 12.0, 1.66, GOLD, 16)
    ribbon(s, "视频按仿真时间 1× 播放 ≠ 仿真已实时；材料未越界 ≠ 工业无损认证。")
    notes(s, "源代码有单元测试、台架、机器人多种子与错误动作检查等分层入口。但最后一轮 S9 全量长测曾按用户要求暂停，不以视频或不同源码的历史结果合并宣称当前完整验收通过。四条线束/卡扣录像用3585854，三条铣削录像用a557bb0。")

    s = new_slide("走向 BMS 拆解：复用机制，补齐真实工艺", "APPLICATION / BMS DISASSEMBLY", status="扩展路线 / 尚非完整示范", subtitle="V3 的 BMS 示范占位页，现改为从现有机制通向真实场景的实施路径。")
    flow(s, [("连接器解锁", "复用卡扣按压与拔出"), ("线束处置", "复用张力与放置审计"), ("紧固 / 粘连", "新建螺纹与胶粘模型"), ("模组移除", "验证支撑、无损与安全")], 2.52, w=2.62, gap=0.50)
    card(s, "可以直接继承", "对象组合、菜谱定义与回归入口\n接触力 / 张力 / 退出与投放逻辑\n轨迹回放与失败证据", 0.66, 4.10, 5.80, 2.10, TEAL, 17)
    card(s, "需要工业方共同提供", "脱敏 CAD、装配关系与工艺步骤\n卡扣 / 螺纹 / 胶粘的实测曲线\n电气隔离、禁切区域与工件保护约束", 6.86, 4.10, 5.80, 2.10, GOLD, 17)
    ribbon(s, "当前没有完整动力电池 BMS 拆解验证；通用铣削演示不代表允许切割电池或带电系统。")
    notes(s, "原 V3 第11页为彭颜龙 BMS 拆解示范占位。本页不伪造电池视频，将现有卡扣/线束能力与将来的BMS流程做映射。电池电气/热安全未建模，不能以本演示指导带电系统作业。")

    s = new_slide("精密装配扩展：谐波减速器作为候选工艺", "APPLICATION / PRECISION ASSEMBLY", status="工艺设计 / 待建模与验证", subtitle="保留 V3 的双臂分工思路；设备精度、实时周期和成品指标需由实测与选型支撑。")
    card(s, "左臂：主操作", "抓取、对准、柔轮预贴合\n波发生器压入与紧固件锁附\n感知引导与过程力监测", 0.66, 2.48, 5.80, 2.11, BLUE, 18)
    card(s, "右臂：支撑与反力", "工件固定与姿态保持\n跟随工艺提供受控支撑\n约束双臂相对运动与内力", 6.86, 2.48, 5.80, 2.11, TEAL, 18)
    flow(s, [("来料与定位", "几何 / 公差 / 夹持"), ("对准与压入", "力—位移接触验证"), ("锁附与润滑", "螺纹 / 摩擦模型"), ("终检与追溯", "质量指标与证据")], 5.04, w=2.62, gap=0.50)
    ribbon(s, "不是已交付场景；不把 ±0.02 mm、±0.5 N、1 ms、±0.005 mm 等未验证数值写为平台实绩。")
    notes(s, "原V3第12页工程架构包含机器人重复精度、力控精度、EtherCAT周期及成品轴向精度等数值，未见型号选型/实测依据。本页保留候选场景与双臂职责，把未佐证数字移出事实陈述；具体指标由后续需求与验证确定。")

    s = preserve(13, "套件矩阵 / 建设目标")
    replace(s, "08 / BENCHMARK SUITES", "SUITES / TARGET MATRIX")
    replace(s, "场景套件把“算法能力”与“工业问题”分开组织，允许团队从基础能力进入，再逐步挑战真实工艺。", "当前原型覆盖线束、卡扣拆解和工具加工；完整套件、数据规模与算法对标仍按阶段建设。")
    notes(s, "套件矩阵继承建设方案，不意味着每个套件已经交付。当前七段演示不等于Spatial/Assembly/Reman的全部种子场景已完成。Tooling 是建设方案架构中列出的拓展方向。")

    s = new_slide("数据与工具链：当前证据轨迹，后续训练数据", "TOOLCHAIN / NOW AND NEXT", status="现有工具 + 后续规划", subtitle="避免把“能保存回放轨迹”说成“已实现遥操作示教与标准训练数据集”。")
    card(s, "现在：验证与审阅", "场景校验 / 仿真 CLI / 脚本基线\nNPZ：物理状态 + info；加工另存材料事件\nJSON：身份、seed、指标与失败证据\n多视角回放与独立 MP4 导出", 0.66, 2.49, 5.80, 3.27, TEAL, 17)
    card(s, "后续：采集与训练", "双 SpaceMouse / VR / 外骨骼遥操作\nRGB-D / 手眼 / 力觉等同步采集\nLeRobot / HDF5 / Zarr 标准化导出\n学习策略训练、数据协议与榜单", 6.86, 2.49, 5.80, 3.27, GOLD, 17)
    ribbon(s, "先稳固材料、接触与回归，再建设采集接口；当前 NPZ 不等价于训练 checkpoint 或专家示教集。")
    notes(s, "替代V3第14页将ibero-record/ibero-cli与标准存储描述为现有功能的表述。旧Trace与StockTrace安全/格式能力有差异，均不是任意外部NPZ的安全沙箱；不在本页夸大统一协议已经完成。")

    s = new_slide("共建方式：先交付一个可验证的工业对象", "CONTRIBUTION / PRACTICAL ENTRY", status="当前参与路径", subtitle="用现有目录和检查工具开始，不要求贡献者等待训练接口或从零搭建仿真底座。")
    flow(s, [("提供对象", "脱敏 CAD + 参数来源"), ("编写菜谱", "配置 / 约束 / task_spec"), ("验证因果", "正例 + 反例 + 变体"), ("提交证据", "源码 / 报告 / 轨迹 / 视频")], 2.56, w=2.62, gap=0.50)
    card(s, "高校 / 算法团队", "先选一个材料或接触机制，形成可重复实验；\n在基线稳定后，再扩展感知、控制与学习策略。", 0.66, 4.19, 5.80, 1.78, BLUE, 17)
    card(s, "工业 / 工艺团队", "提供参数测量、误差容限与失败定义；\n共同确认哪些近似足以支持目标工艺判断。", 6.86, 4.19, 5.80, 1.78, TEAL, 17)
    ribbon(s, "30–50 条专家轨迹是后续数据共建目标，不是当前 Core 材料/机构贡献的强制前置。")
    notes(s, "V3的ibero create-scene / ibero teleop / ibero check属于建设方案目标命令，不在此伪装为现行统一CLI。当前命令见README：python -m ibero.check_scene、verify_industrial等。")

    s = new_slide("指标：任务完成之外，还要看约束与证据", "METRICS / IMPLEMENTATION BOUNDARIES", subtitle="继承建设方案的五类评价目标，同时标清 Core-0.1 当前能够报告的内容。")
    widths = [2.15, 5.13, 4.72]
    rows = [
        ("评价维度", "当前可观测 / 可检查", "不能直接外推"),
        ("任务完成", "接触交接、拔出放置、目标孔槽判定", "完整工厂工艺已交付"),
        ("工件保护", "张力 / 力 / 形变 / 侵入等阈值", "真实塑性损伤或寿命已预测"),
        ("力与内力", "张力、接触力、切削载荷及越界记录", "通用双臂内力控制已完备"),
        ("碰撞安全", "禁止接触与机器人姿态等运行审计", "实时规避与形式化安全保证"),
        ("泛化与稳健性", "规定 seed / 参数变体 / 反事实检查", "大规模 OOD 或实机迁移已通过"),
    ]
    y = 2.47
    for r, values in enumerate(rows):
        x = 0.66
        for w, value in zip(widths, values):
            box(s, x, y, w - 0.025, 0.55, "16465E" if r == 0 else (CARD if r % 2 else "0D2B3C"))
            text(s, value, x + 0.14, y + 0.11, w - 0.28, 0.38, 13.2, TEAL if r == 0 else WHITE, r == 0)
            x += w
        y += 0.58
    ribbon(s, "失败与模型适用边界也必须记录；不能只汇报成功视频或一个总 passed。")
    notes(s, "将V3指标体系落实到当前字段，但不把报警代理当真实损伤标签。少样本泛化与训练基线尚未铺开。所有结果应引用明确版本、任务、seed集合与分母。")

    s = new_slide("当前交付：机制原型已可演示，平台仍在迭代", "RELEASE / CORE-0.1 STATUS", subtitle="把实际成果与建设计划分开陈述，给下一轮开发一个清晰起点。")
    card(s, "已有可运行原型", "线束交接 / 双端张力控制\n弹性卡扣拔出 / 带线束拔出放置\n通孔 / 直槽 / 浅腔三种加工目标\n脚本基线、检查入口、证据回放", 0.66, 2.52, 5.80, 3.31, TEAL, 18)
    card(s, "后续按门槛推进", "实测物性与工艺系数标定\n性能优化及最终同源完整回归\n扩展装配 / 螺纹 / 胶粘等机制\n遥操作、学习策略与规模化基准", 6.86, 2.52, 5.80, 3.31, GOLD, 18)
    ribbon(s, "当前基线是脚本控制器；ACT / Diffusion Policy / OpenVLA 等仍为后续对标目标。")
    notes(s, "替换原V3第17页把Spatial、Assembly、Disassembly、Reman四种种子场景与主流学习模型排成既成首发交付的表述。建设目标仍保留，但现阶段不虚构这些任务或训练结果。")

    s = preserve(18, "原方案时间目标 / 非交付承诺")
    replace(s, "13 / ROADMAP", "ROADMAP / GATED DELIVERY")
    replace(s, "12个月路线图：从底座可用，到社区共建，再到开源赛事", "建设路线：先物理可信，再数据与社区，最后规模化")
    replace(s, "路线图把平台建设拆成三个递进阶段，每阶段都有可被社区使用或贡献的结果。", "原方案的 1–3 / 4–6 / 7–12 月作为规划参考；实际推进以标定、完整回归和资源到位为门槛。")
    replace(s, "ibero-record 与数据标准", "物性标定 / 完整回归 / 数据预研")
    replace(s, "4个种子场景 + Baseline", "现有机制稳定 → 扩展种子场景")
    replace(s, "10+垂直细分场景", "逐步扩展场景、采集与 Baseline")
    replace(s, "正式开源与官方榜单", "确认资产/代码许可后开放发布")
    replace(s, "IROS / ICRA / CoRL 挑战赛", "论文、榜单与挑战赛（拟）")
    notes(s, "时间段为原建设方案相对时间规划，不是从今日开始的硬性承诺，也不表示对应阶段已完成。先关闭当前Core的标定与S9完整回归门槛，随后考虑更大版本。官方赛事未获确认，不宣称已经与会议达成合作。项目代码开放发布还需明确许可证和资产分发边界。")
    s = preserve(19)
    replace(s, "14 / INVITATION", "INVITATION / CO-BUILD IBERO")
    replace(s, "可复现数据、公开榜单、Benchmark论文与工业验证", "可复现证据、实测标定，逐步走向数据、论文与工业验证")
    notes(s, "邀请高校与企业贡献真实对象、测量曲线与工艺失败定义。V4为工程原型演示，不将总体愿景等同当前成果。展示视频建议主讲线束交接、卡扣拔出、通孔三段；其余按讨论需要播放，全部七段合计约4分钟。")

    # 用同一演讲稿内部关系重新排序；未使用的占位页断开关系，不残留虚假内容。
    ordered_ids = [slide.slide_id for slide in order]
    wanted = set(ordered_ids)
    ids = prs.slides._sldIdLst
    by_id = {int(el.id): el for el in ids}
    for el in list(ids):
        if int(el.id) not in wanted:
            prs.part.drop_rel(el.rId)
        ids.remove(el)
    for n, (slide, slide_id) in enumerate(zip(order, ordered_ids), 1):
        ids.append(by_id[slide_id])
        for sh in list(slide.shapes):
            if sh.has_text_frame and sh.text.strip().isdigit() and sh.left > Inches(12) and sh.top < Inches(1):
                sh._element.getparent().remove(sh._element)
        if n > 1:
            text(slide, f"{n:02d}", 12.47, 0.31, 0.3, 0.23, 10, MUTED)
    prs.core_properties.title = "IBERO 工业级双臂具身智能操作基准平台 — V4"
    prs.core_properties.subject = "建设愿景、Core-0.1 技术实现与七段内嵌演示"
    prs.core_properties.comments = "保留 V3 原文件；视频源版本与限制见各页备注。"
    prs.save(OUTPUT)
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Original V3 unexpectedly changed")
    with ZipFile(OUTPUT) as archive:
        videos = [name for name in archive.namelist() if name.startswith("ppt/media/") and name.endswith(".mp4")]
        if len(videos) != 7:
            raise RuntimeError(f"Expected 7 embedded videos, got {len(videos)}")
        expected = {entry["sha256"] for entry in entries if entry["kind"] == "video"}
        actual = {hashlib.sha256(archive.read(name)).hexdigest() for name in videos}
        if expected != actual:
            raise RuntimeError("Embedded video bytes do not match verified originals")
    manifest = {"source_v3_sha256": source_hash, "output": OUTPUT.name,
                "output_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                "slide_count": len(order), "embedded_videos": len(videos), "slides": entries}
    (OUT / "build_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"file": str(OUTPUT), "slides": len(order), "embedded_videos": len(videos), "bytes": OUTPUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    # 用户已撤销全稿校订；当前默认入口必须完整保留原有页面。
    from revise_ibero_v4_additive import main as additive_main
    additive_main()
